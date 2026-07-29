from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\实验\CARD")
OUT = ROOT / "docs" / "RSICC_Latest_Literature_LEVIR_MCI_SECOND_CC_CARD_Review_zh_2026-07-26.docx"

# Design preset: standard_business_brief.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B7A"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_AMBER = "FFF7E6"
PALE_RED = "FDECEC"
GRID = "C8D0D9"
WHITE = "FFFFFF"
ASCII_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"


def set_run_font(run, size=None, bold=None, color=None, italic=None, name=ASCII_FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        old_grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True, widow=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))
    if widow:
        p_pr.append(OxmlElement("w:widowControl"))


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    paragraph.add_run(" 页")


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), ASCII_FONT)
    r_fonts.set(qn("w:hAnsi"), ASCII_FONT)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "21")
    r_pr.append(sz)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def paragraph_border_and_shading(paragraph, fill, left_color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), left_color)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    ind = paragraph.paragraph_format
    ind.left_indent = Inches(0.16)
    ind.right_indent = Inches(0.10)
    ind.space_before = Pt(5)
    ind.space_after = Pt(9)
    ind.line_spacing = 1.10


def body(doc, text="", bold_lead=None, color=None, italic=False):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True, color=color)
        set_run_font(p.add_run(text[len(bold_lead):]), color=color, italic=italic)
    else:
        set_run_font(p.add_run(text), color=color, italic=italic)
    set_paragraph_keep(p, keep_lines=False)
    return p


def lead_callout(doc, label, text, fill=PALE_BLUE, color=INK, left_color=BLUE):
    p = doc.add_paragraph()
    set_run_font(p.add_run(label + "  "), bold=True, size=11, color=left_color)
    set_run_font(p.add_run(text), size=10.5, color=color)
    paragraph_border_and_shading(p, fill, left_color)
    return p


def flow_box(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    set_run_font(p.add_run(text), size=9.5, bold=True, color=DARK_BLUE)
    paragraph_border_and_shading(p, LIGHT, DARK_BLUE)
    return p


def create_numbering(doc, kind="bullet"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def list_item(doc, list_kind, text, lead=None):
    # Word's built-in List Bullet/List Number styles carry real numbering
    # definitions and are more portable across Word/LibreOffice renderers.
    style_name = "List Bullet" if list_kind == "bullet" else "List Number"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if lead and text.startswith(lead):
        set_run_font(p.add_run(lead), bold=True)
        set_run_font(p.add_run(text[len(lead):]))
    else:
        set_run_font(p.add_run(text))
    set_paragraph_keep(p, keep_lines=False)
    return p


def add_link_line(doc, links):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    for idx, (label, url) in enumerate(links):
        if idx:
            set_run_font(p.add_run("  |  "), size=10, color=MUTED)
        add_hyperlink(p, label, url)
    return p


def add_reference(doc, number, citation, links):
    """Add one compact bibliography entry with clickable primary-source links."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    set_run_font(p.add_run(f"[{number}] "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(citation))
    if links:
        set_run_font(p.add_run(" "), color=MUTED)
    for idx, (label, url) in enumerate(links):
        if idx:
            set_run_font(p.add_run(" · "), color=MUTED)
        add_hyperlink(p, label, url)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_paragraph_keep(p, keep_next=True, keep_lines=True)
    return p


def add_table(doc, headers, rows, widths, font_size=8.8, header_fill=LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(text), size=font_size, bold=True, color=INK)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            if isinstance(value, tuple):
                text, bold = value
                set_run_font(p.add_run(text), size=font_size, bold=bold)
            else:
                set_run_font(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(4)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = ASCII_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    body_style = styles["Body Text"]
    body_style.font.name = ASCII_FONT
    body_style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    body_style.font.size = Pt(11)
    body_style.paragraph_format.space_after = Pt(6)
    body_style.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = ASCII_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    table_style = styles["Table Grid"]
    table_style.font.name = ASCII_FONT
    table_style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    table_style.font.size = Pt(8.8)


def configure_section(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("RSICC  |  最新文献·数据集·CARD 方法审计"), size=8.5, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    for run in p.runs:
        run.clear()
    add_page_field(p)
    for run in p.runs:
        set_run_font(run, size=8.5, color=MUTED)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    set_run_font(p.add_run("RESEARCH REVIEW  •  2026-07-26"), size=9, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("遥感图像变化描述最新文献综述"), size=23, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("兼论 LEVIR-MCI / SECOND-CC 与当前 CARD 实验"), size=16.5, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("副标题  "), size=10, bold=True, color=MUTED)
    set_run_font(p.add_run("2021—2026 方法演进、最新论文、数据集溯源、代码对比与实验路线"), size=10.5, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("代码快照  "), size=10, bold=True, color=MUTED)
    set_run_font(p.add_run("branch recode · commit c906eb3 · 本地只读审计"), size=10.5, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("证据范围  "), size=10, bold=True, color=MUTED)
    set_run_font(p.add_run("检索截至 2026-07-26；论文原文、出版社、官方仓库与本地源码/配置"), size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    bullet_id = "bullet"
    number_id = "number"

    props = doc.core_properties
    props.title = "遥感图像变化描述最新文献综述：LEVIR-MCI、SECOND-CC 与 CARD 对比"
    props.subject = "截至 2026-07-26 的 RSICC 方法演进、最新论文、数据集与代码审计"
    props.author = "Codex"
    props.keywords = "RSICC, remote sensing change captioning, LEVIR-MCI, SECOND-CC, Change-Agent, MModalCC, CARD, 2026"

    add_title_block(doc)
    lead_callout(
        doc,
        "结论先行",
        "截至 2026-07-26，检索到的最新直接 RSICC 预印本是 2026-07-03 提交的 LBTCap；2026 年研究正从单一差异注意力转向变化检测先验、层次路由、频域鲁棒性、VLM/LLM 后训练和新型多任务数据。你的两个最新数据集可确认是 LEVIR-MCI 与 SECOND-CC（本地使用 SECOND-CC-AUG）。当前实现是“以 CARD 为统一主干的辅助监督/语义融合扩展”，不是 Change-Agent/MCINet 或 MModalCC 的逐模块复现；且在修复验证随机解码、SECOND 类别口径、真值语义图测试与缺失工件前，还不能形成可审稿的性能结论。",
    )

    heading(doc, "执行摘要", 1)
    for text, lead in (
        ("数据判断：LEVIR-MCI 的本地规模与论文一致；SECOND-CC 本地目录是论文的 AUG 协议版本，训练集与验证集各自扩增一倍、测试集保持原始。", "数据判断："),
        ("方法判断：LEVIR-MCI 代码把三类变化掩膜与 caption 派生的 21 维语义标签作为训练期辅助监督；SECOND-CC 代码把真值前后语义图直接送入单层语义 cross-attention，并附加 14×14 语义差异预测。", "方法判断："),
        ("对比判断：论文方法都使用端到端或可微的专门多分支视觉编码；当前代码以冻结的离线 ResNet101 stage-3 特征为输入，保留 CARD 的公共/差异表征。这更适合做“在固定 CARD 主干上验证监督信号价值”的消融研究，不适合宣称复现或替代原论文架构。", "对比判断："),
        ("优先动作：先统一验证/测试解码、修正 SECOND 类别定义、生成完整特征并建立 validation-only 锁定清单；再做多随机种子与官方基线对照。", "优先动作："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "报告范围与判定口径", 2)
    body(doc, "本报告把“论文方法是否先进”“当前代码是否忠实复现”“当前实验是否可比较”拆成三个独立问题。论文表格中的公开结果用于理解方法作用，不作为本地模型已达到该结果的证据。本地历史 Markdown 中曾记录过若干单次 Test 数字，但当前工作区没有对应 checkpoint，正式实验汇总也明确标记为 missing/failed，因此本文不据此宣称性能提升。")

    heading(doc, "A. 文献检索范围与最新结论", 1)
    lead_callout(
        doc,
        "最新边界",
        "按“论文核心任务直接生成双时相遥感变化描述”这一严格口径，截至 2026-07-26 可核验的最新论文是 LBTCap（arXiv:2607.03320，2026-07-03）。6 月下旬还出现 JL1-CC&QA、RSICCLLM 与 DFM。预印本的模型、数据和数值可能在正式发表前变化，本文把它们与已正式发表论文分开标注。",
        fill=PALE_AMBER,
        left_color="C58A00",
    )
    body(doc, "检索优先使用论文原文、DOI/出版社页、会议论文集与作者官方仓库。纳入标准包括：直接 RSICC 模型、为 RSICC 新增数据/评测的工作，以及对当前 LEVIR-MCI、SECOND-CC 或 CARD 设计有直接方法启示的联合变化检测/描述论文。仅把“变化检测”或普通单图遥感描述作为辅助背景，不混入主表。")
    add_table(
        doc,
        ["证据等级", "含义", "本文处理"],
        [
            ["正式发表", "DOI/出版社或会议论文集可核验", "可作为稳定方法来源；结果仍受数据、分词和输入协议限制"],
            ["录用/在线优先", "作者或出版社给出录用/online-first，但卷期可能尚未生效", "注明在线状态，不把未来卷期日期当作检索日期"],
            ["预印本", "arXiv 原文可核验，可能尚未同行评审", "用于识别最新方向；所有指标标为作者报告，不视作确定 SOTA"],
        ],
        [1550, 3300, 4510],
        font_size=8.5,
    )

    heading(doc, "A.1 任务与方法演进：从 CNN–RNN 到变化先验和 VLM 后训练", 2)
    add_table(
        doc,
        ["阶段", "代表工作", "方法变化与仍未解决的问题"],
        [
            ["2021–2022\n任务形成", "Chouaf et al.; Hoxha et al.; RSICCFormer / LEVIR-CC", "由 CNN–RNN/SVM 概念验证转向双分支 Transformer 和大规模标准数据；仍依赖单域、建筑变化为主的 LEVIR-CC。"],
            ["2023\n显式差异", "PSNet; PromptCC; Chg2Cap", "多尺度差异、层次跨时相注意与 change/no-change 解耦成为主线；参数量大、跨区域验证不足。"],
            ["2024\n高效与多任务", "RSCaMa; CARD; Change-Agent; Semantic-CC", "SSM 降低视觉成本，CARD 分离公共/差异上下文，检测掩膜、SAM/Vicuna 与 Agent 被引入；任务标签不一致和小目标仍突出。"],
            ["2025\n数据与基础模型", "SECOND-CC; KCFI; SAT-Cap; RSCC; ChangeIMTI", "开始重视真实扰动、多灾种/多任务数据、检测约束与指令微调；oracle 语义输入、合成文本和事实幻觉成为新风险。"],
            ["2026\n先验、路由与后训练", "STAND; HiSem; DFM; RSICCLLM; LBTCap", "变化检测先验、频域鲁棒、层次 MoE、文本语义对齐、SFT/偏好优化和实时轻量化并行；统一基准与事实级评测仍缺失。"],
        ],
        [1450, 2550, 5360],
        font_size=8.15,
    )

    heading(doc, "A.2 2025—2026 重点论文速览", 2)
    add_table(
        doc,
        ["编号 / 状态", "论文与时间", "方法核心", "结果、价值与限制"],
        [
            ["[R1]\n预印本", "LBTCap\n2026-07-03", "截断 ResNet101 + 双边共享 Q/K 注意；拼接 value、可学习双边权重、GQA；单层编码/解码。", "作者报 39.99M 参数、编码器 2.78M，batch-1 约 18.94–62.26 FPS；重在速度/精度折中，不是绝对精度第一。"],
            ["[R2]\n预印本", "JL1-CC&QA\n2026-06-30", "为 JL1-CD 增加 caption 与 QA；MLLM 生成、视觉落地 LLM 审核、专家复核三阶段。", "5,000 对、17,021 captions、20,060 QA；是数据/评测贡献，不应当作新模型排行榜结果。"],
            ["[R3]\nECCV 2026 录用", "RSICCLLM\n2026-06-26", "Qwen2.5-VL-7B；差异感知 SFT + 可微配准/中心差分/Hough 特征；双负例偏好优化并加 reverse-KL。", "构建约 35k 训练和约 5k OOD 指令/偏好样本；自建基准与经典 LEVIR-CC n-gram 表不可直接横比。"],
            ["[R4]\nICME 2026 录用", "DFM\n2026-06-25", "ViT 与冻结变化检测编码器；联合差异建模读取 CD 先验；冻结句向量模型，用文本引导门控对比损失。", "作者报 LEVIR-CC B4 66.26 / CIDEr-D 142.51；对比损失排除 no-change，避免所有“不变”句塌缩。"],
            ["[R5]\nTGRS 2026", "HiSem\n2026-05-14", "方向感知跨时相差异注意；先路由 change/no-change，再让 change token 进入分组约束 MoE；课程训练。", "作者报 LEVIR-CC B4 65.82 / CIDEr-D 138.86；路由错误会级联，专家收益需与参数/吞吐同时报告。"],
            ["[R6]\n预印本", "PTNet + UCCD\n2026-05-06", "CLIP ViT-L/14 LoRA；掩膜差异聚类初始化原型；检测/描述头级门控；Qwen2-1.5B LoRA 与文本对比对齐。", "UCCD 9,000 对、45,000 captions 由 5 个 VLM 生成再人工核验；截至检索日数据下载仍待开放。"],
            ["[R7]\n预印本", "STAND\n2026-04-25", "before + 预训练 CD mask + after 组成短视频；伪 after 的双向对比；频域高通小变化注意与实体类别先验。", "证明外部 CD 先验可作为软条件，而非硬删区域；部署仍依赖 CD 模型质量。"],
            ["[R8]\n正式发表", "Disturbance-Robust RSCC\n2026-03-03", "多频表征分离前景/背景；前景做跨时相注意，低频相关自监督对齐旋转背景。", "直接针对旋转等非语义扰动；构造扰动集上的提升不等同真实跨传感器泛化。"],
            ["[R9]\n正式发表", "Difference Semantic Prior Guidance\n2026", "Siamese ResNet101 多尺度；深层对称差异上下文 + 浅层辅助差异；语义提示交叉精炼解码。", "强调浅/深差异互补；论文同时承认参数和部署时延更高。"],
            ["[R10]\n正式发表", "MVLT-LoRA-CC\n2026", "冻结 ViT，多时相 token 加时间索引；多模态 LLM cross-attention；只对适配/多模态模块做 LoRA。", "给出 Complementary Consistency Score 以拆分变化描述与 no-change 判断；仍需事实级人工核验。"],
            ["[R11]\n在线优先", "Ricci et al.\n2026-06", "Vicuna-13B / Otter-9B 三种零样本策略；GPT-3.5 扩写 SECOND；提出 LLM 事实匹配 FMScore。", "显示零样本 VLM 的 n-gram 不强但事实分数可竞争；合成参考与 LLM judge 都可能引入同源偏差。"],
            ["[R12]\n正式发表", "KCFI\nTIP 2025", "ViT + key feature perceiver；变化检测 decoder 约束关键特征；指令微调 LLM；动态任务权重。", "检测约束提升关键变化感知，但模型大、训练复杂；与当前低分辨率辅助头不是同等监督强度。"],
            ["[R13]\n预印本", "SAT-Cap\n2025-01-14", "空间–通道注意编码器；余弦相似度差异引导融合；单阶段 caption decoder。", "作者报 LEVIR-CC CIDEr 140.23；结构相对简洁，但核验时未见官方代码。"],
            ["[R14]\nNeurIPS 2025 D&B", "RSCC 灾害数据集\n2025", "31 个灾害事件、62,351 对；QvQ-Max 生成、Qwen2.5-Max 校正，再抽样人工验收。", "显著扩展灾害域；机器生成/校正链可能产生统一措辞和事件先验泄漏，应做跨事件划分。"],
            ["[R15]\n预印本", "ChangeIMTI / ChangeVG\n2025-09", "同一数据支持 caption、二分类、计数和定位；细粒度空间 + 高层语义提示 Qwen2.5-VL-7B。", "适合多任务事实性分析；自报增益需等正式协议和公开实现复核。"],
            ["[R16]\nTGRS 2026", "SAM-guided region mining\n2025-11 预印本", "CNN/Transformer 全局特征 + SAM 语义/运动变化区域 + 知识图谱；cross-attention 解码。", "已正式发表；区域先验丰富，但 SAM 区域、运动线索与知识图谱的错误会共同传导。"],
            ["[R17]\nTIP 2026", "LCCC\n2026", "构造 CD↔CC critique-and-correction 闭环：检测与描述互查错误、生成 corrective guidance，再联合精炼 mask/caption。", "与 LEVIR-MCI 最直接相关；需 mask+caption 联合监督，存在互相放大错误的风险；WHU 只证明检测泛化。"],
            ["[R18]\nTGRS 2026", "Weakly Paired LMM\n2026", "用 changed object + action 弱配对 tags 替代完整句配对；视觉变化感知与语言解释两阶段；同义扩写回齐。", "论文称可达全监督 SOTA 相当表现而非明确超越；tags 会丢数量、方向和关系；未见公开代码。"],
            ["[R19]\n正式发表", "DeltaVLM\n2026-02-08", "选择性微调双时相视觉编码；IDPM/CSRM 过滤无关变化；instruction-guided Q-former + 冻结 Vicuna-7B。", "ChangeChat-105k 支持六类交互任务；B4 62.51 / CIDEr 136.72 来自自建测试集，不是 LEVIR-CC 主表。"],
        ],
        [850, 1730, 3420, 3360],
        font_size=7.55,
    )

    heading(doc, "A.3 最新方法论的四条主线", 2)
    heading(doc, "A.3.1 更强差异并不只等于更深注意力", 3)
    body(doc, "LBTCap 以共享 Q/K、双边 value 和 grouped-query attention 证明，跨时相交互可以通过结构共享而不是堆叠层数来提速；DFM 把变化检测编码器当成冻结先验，并让图像差异与整句语义对齐；HiSem 把“是否变化”和“如何描述变化”拆成路由与专家；LCCC 则让 CD 与 CC 互相 critique/correct，形成真正的任务闭环。共同结论是：下一步重点是差异表示的条件化、选择性与跨任务纠错，而非单纯增加 Transformer 深度。")
    body(doc, "对当前 CARD 的直接启示是保留公共/差异分解，但增加可解释的条件：用 predicted CD prior 作为软注意、用 change/no-change 头控制 decoder 路径，并加入文本级对比约束。DFM 排除 no-change 对比样本尤其重要，因为大量相同“不变”参考会让语义空间出现伪聚类。")

    heading(doc, "A.3.2 鲁棒性从增强数据转向显式扰动建模", 3)
    body(doc, "SECOND-CC 首先把错位、模糊、亮度与尺度差异纳入数据；STAND 将频域小变化线索和全局上下文分开建模；Disturbance-Robust RSCC 用低频自监督对齐旋转背景。与仅做离线旋转/亮度增强相比，这些方法都显式约束“扰动不应被描述成语义变化”。因此鲁棒性实验必须构造可控扰动强度曲线，并同时报告 caption 事实性和 no-change 误报率。")

    heading(doc, "A.3.3 基础模型带来知识，也带来新的证据风险", 3)
    body(doc, "KCFI、MVLT-LoRA-CC、DeltaVLM、ChangeVG 和 RSICCLLM 分别通过检测约束、LoRA、指令引导、多任务指令和偏好优化把 VLM/LLM 引入 RSICC。Weakly Paired LMM 进一步表明完整句配对不是唯一监督形式，但 object/action tags 会成为信息瓶颈。基础模型改善语言流畅度和开放词汇能力，也更容易把常识当成影像证据。Ricci 等提出 FMScore 的动机正是 n-gram 无法判断事实；不过 LLM 评审也受模型偏好影响，必须和对象/动作/方向/数量的规则指标及人工双盲评审组合使用。")

    heading(doc, "A.3.4 新数据集正在改变“泛化”的定义", 3)
    body(doc, "UCCD 强调超高分辨率短时隔城市施工，RSCC 强调跨灾害事件，JL1-CC&QA 与 ChangeIMTI 强调 caption 之外的 QA、计数和定位。未来的“泛化”不能只指 LEVIR-CC 随机测试集；更可信的方案是按城市/事件留出、跨传感器、跨分辨率，并明确区分人工文本、VLM 生成文本和专家复核文本。")

    heading(doc, "A.3.5 与当前项目最直接的三篇新增正式论文", 3)
    body(doc, "LCCC（TIP 2026）与当前 LEVIR-MCI 分支的可比性最高：两者都同时拥有 mask 和 caption，但 LCCC 让 CD/CC 预测互相批评并生成纠正信号，当前代码只有辅助损失对 CARD 表征的单向塑形。它应被列为强外部基线；如果移植，优先从单轮、可关闭的 soft correction 开始，并记录一次纠错前后的 mask/caption 错误是否同步下降，避免闭环放大错误。")
    body(doc, "Weakly Paired LMM（TGRS 2026）把完整句监督压缩成 object/action tags，这与当前 21 维 caption-derived 标签表面相似，但研究条件不同：当前代码仍用完整 caption 训练语言交叉熵，tags 只是附加头，不能据此声称“弱配对训练”。若要做公平复现，应在视觉阶段隐藏完整句，只保留独立 tags，再评估数量、方向和多对象关系的损失。")
    body(doc, "DeltaVLM（Remote Sensing 2026）把任务扩展到交互式二分类、计数、定位、QA 和多轮对话。它的 105,107 条 ChangeChat 样本由 LEVIR-CC/MCI 派生并含规则/GPT 辅助；caption B4 62.51、CIDEr 136.72 均来自自建测试集，不能搬到标准 LEVIR-CC 排名。对当前项目，它适合作为长期 VLM 分支和事实性任务设计参考，而不是短期替换 CARD 的同协议 baseline。")

    heading(doc, "A.4 最新文献与当前 CARD 实验的直接映射", 2)
    add_table(
        doc,
        ["最新趋势", "代表论文", "当前代码状态", "最有价值的改进"],
        [
            ["轻量跨时相交互", "LBTCap", "冻结 14×14 ResNet 特征；CARD 跨时相注意；尚未报告参数/FLOPs/FPS", "以相同 decoder 对比共享 Q/K 或 GQA，补参数量、显存、batch-1 延迟"],
            ["变化检测先验", "DFM, STAND, KCFI", "LEVIR 仅把 14×14 mask 当辅助损失；SECOND 直接使用 GT 语义图", "把 GT 输入替换为冻结/可训练 predicted prior，并做 prior 质量退化曲线"],
            ["检测—描述闭环纠错", "LCCC", "mask/semantic 头只提供单向辅助梯度；没有让 caption 反查 mask 错误", "把 LCCC 作为 LEVIR-MCI 强基线；先做单轮软纠错，监控两任务错误传播"],
            ["change/no-change 层次路由", "HiSem, PromptCC", "两类样本共用同一解码路径，只在评测时分组", "增加软路由/门控，报告路由混淆矩阵与错误级联"],
            ["弱标签与文本语义对齐", "DFM, Weakly Paired LMM, RSICCLLM", "LEVIR 的 21 维标签由关键词规则从 caption 派生", "加入句向量对比；把 21 维 tags 当可控弱监督而非真值；专评数量/方位/关系丢失"],
            ["扰动鲁棒性", "SECOND-CC, STAND, Disturbance-Robust", "SECOND 使用静态 AUG；视觉特征离线，训练时不能在线施扰", "建立旋转/错位/模糊/亮度强度矩阵；必要时改为在线或多版本特征"],
            ["VLM/LLM 与交互", "MVLT-LoRA, DeltaVLM, ChangeVG, RSICCLLM", "当前为 2 层 Transformer decoder，无开放词汇后训练或多轮任务", "作为后续高算力分支；DeltaVLM 仅在 ChangeChat-105k 比较，不能搬用为 LEVIR-CC 基线"],
            ["事实级评测", "Ricci FMScore, JL1-CC&QA", "已有 BLEU/METEOR/ROUGE/CIDEr/SPICE 与分组工具", "新增对象、动作、方向、数量、遗漏/幻觉；BERTScore/SBERT/LLM 评审只作补充"],
        ],
        [1800, 1900, 2700, 2960],
        font_size=8.0,
    )
    lead_callout(doc, "研究判断", "与 2026 年前沿最契合、且不需要立即换成大模型的路线是：CARD 公共/差异分解 + predicted change prior + no-change-aware 文本对比 + 鲁棒扰动课程；LCCC 作为联合 CD/CC 的强外部基线，并以单轮软纠错做扩展消融。它能保留当前工程资产，同时解决 oracle 语义、低分辨率和事实性三项核心问题。")

    heading(doc, "1. 两个最新数据集的确认与版本溯源", 1)
    body(doc, "根据目录时间、Git 提交与配置新增记录，两个最新加入当前仓库的数据集是 LEVIR-MCI 和 SECOND-CC-AUG。提交 f6a8c1f（2026-06-22）首次同时增加两套配置与数据加载适配，a2f9605（2026-06-24）继续补齐论文实验流程；当前快照为 c906eb3。")

    add_table(
        doc,
        ["项目", "LEVIR-MCI", "SECOND-CC（论文原版）", "本地 SECOND-CC-AUG", "比较含义"],
        [
            ["图像对", "10,077", "6,041", "10,855", "本地 SECOND 不是原版规模"],
            ["划分", "6,815 / 1,333 / 1,929", "4,219 / 595 / 1,227", "8,438 / 1,190 / 1,227", "train/val 各扩增一倍；test 不扩增"],
            ["描述", "50,385（每对 5 条）", "30,205（每对 5 条）", "54,197 条本地条目", "53 个本地样本少于 5 条参考，需保留真实计数"],
            ["变化比例", "5,038 change / 5,039 no-change", "4,337 / 1,704", "7,788 / 3,067", "分组指标不可省略"],
            ["额外监督", "3 类变化掩膜：背景/道路/建筑", "前后 6 类土地覆盖语义图", "同原版语义图 + 4,814 个增强样本", "LEVIR 是变化语义；SECOND 是时相语义"],
            ["空间规格", "256×256；0.5 m/pixel", "512×512 切为 256×256；0.5–3 m/pixel", "沿用 256×256 子块", "SECOND 的分辨率/配准扰动更复杂"],
        ],
        [1050, 1750, 1900, 1880, 2780],
        font_size=8.2,
    )

    lead_callout(
        doc,
        "版本边界",
        "文中统一写作“SECOND-CC（本地使用 SECOND-CC-AUG）”。若与论文原表比较，必须同时注明：是否使用 AUG、是否使用真值语义图、验证/测试解码方式，以及 train/val/test 划分。",
        fill=PALE_AMBER,
        left_color="C58A00",
    )

    heading(doc, "1.1 本地数据与论文数据的一致性", 2)
    for text, lead in (
        ("LEVIR-MCI：本地图像对数、划分、每图五条描述及 change/no-change 数量均与官方发布协议一致。", "LEVIR-MCI："),
        ("SECOND-CC-AUG：本地 4,814 个文件名带 random_augment 的样本全部位于 train/val，test 的 1,227 对未扩增；这与论文的公平测试设计一致。", "SECOND-CC-AUG："),
        ("描述条数：本地 SECOND 不是简单的 10,855×5；有 5 个样本含 2 条、15 个含 3 条、33 个含 4 条参考。评测与论文对照时应使用实际 annotation 文件，不能人工补齐或按理论值统计。", "描述条数："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "1.2 直接相关论文与官方入口", 2)
    body(doc, "LEVIR-MCI 来自 Change-Agent 论文；SECOND-CC 与其专用模型 MModalCC 由同一篇论文提出。CARD 则是当前代码的原始主干论文，属于对比所需的第三条方法来源。")
    add_link_line(doc, [
        ("Change-Agent / LEVIR-MCI（arXiv）", "https://arxiv.org/abs/2403.19646"),
        ("DOI", "https://doi.org/10.1109/TGRS.2024.3425815"),
        ("官方代码与数据", "https://github.com/Chen-Yang-Liu/Change-Agent"),
    ])
    add_link_line(doc, [
        ("SECOND-CC / MModalCC（arXiv）", "https://arxiv.org/abs/2501.10075"),
        ("DOI", "https://doi.org/10.1109/JSTARS.2025.3600613"),
        ("官方代码与数据", "https://github.com/ChangeCapsInRS/SecondCC"),
    ])
    add_link_line(doc, [
        ("CARD 主干论文（ACL Anthology）", "https://aclanthology.org/2024.acl-long.430/"),
        ("PDF", "https://aclanthology.org/2024.acl-long.430.pdf"),
    ])

    heading(doc, "2. LEVIR-MCI 论文分析：Change-Agent 与 MCINet", 1)
    body(doc, "论文：Chenyang Liu 等，Change-Agent: Toward Interactive Comprehensive Remote Sensing Change Interpretation and Analysis，IEEE TGRS 2024。论文同时提出 LEVIR-MCI 数据集、一个多层次变化解释模型 MCINet，以及在其上调用工具的 LLM Agent。")
    add_link_line(doc, [
        ("论文原文", "https://arxiv.org/html/2403.19646v3"),
        ("IEEE", "https://ieeexplore.ieee.org/document/10591792"),
        ("官方仓库", "https://github.com/Chen-Yang-Liu/Change-Agent"),
    ])

    heading(doc, "2.1 论文要解决的问题", 2)
    body(doc, "普通变化描述只输出一句文本，变化检测只输出掩膜，两者都不足以回答“哪里变了、变成什么、数量多少、如何交互查询”。Change-Agent 把像素级变化语义和句子级描述放进同一个多任务模型，再由 LLM 将用户问题转为工具调用，从而覆盖分割、描述、计数与组合查询。")

    heading(doc, "2.2 MCINet 方法论拆解", 2)
    flow_box(doc, "双时相影像 → 共享 SegFormer-B1 多尺度骨干 → 双分支 BI3 交互 → 变化掩膜 + 文本描述 → LLM 工具编排")
    for text, lead in (
        ("共享底座、双任务分支：两个时相经过权重共享的 SegFormer-B1；检测分支使用多尺度特征，描述分支强调高层语义。共享底座让像素边界与语言语义互相约束。", "共享底座、双任务分支："),
        ("BI3 交互模块：每个分支堆叠 BI3。LPE 用 3×3、5×1、1×5 卷积补充局部感受野；GDFA 以两时相差分引导全局注意，让模型在配准扰动和大范围变化中仍能定位差异。", "BI3 交互模块："),
        ("检测分支：CBF 将 T1、差异/相似度和 T2 多尺度特征拼接，并自底向上恢复空间细节，输出背景/道路/建筑三类变化掩膜。", "检测分支："),
        ("描述分支：高层双时相特征经 BI3 与卷积投影后送入 Transformer 解码器，自回归生成变化描述。", "描述分支："),
        ("Agent 层：LLM 根据自然语言请求生成并执行 Python 工具调用，把掩膜、描述、计数等能力组合为交互式分析。这一层不是 MCINet 训练本身，也不在当前 CARD 项目中。", "Agent 层："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "2.3 训练、损失与指标", 2)
    body(doc, "像素分支和文本分支都用交叉熵。论文没有使用固定 λ 相加，而是按当前损失自身的 stop-gradient 值归一化：L = Ldet/detach(Ldet) + Lcap/detach(Lcap)，意图让两个任务在不同量纲下保持近似相等的梯度贡献。训练采用 Adam、初始学习率 1e-4、200 epochs；骨干先联合训练，BLEU-4 与 mIoU 长期稳定后冻结骨干，再分别优化分支。主要指标为 mIoU、BLEU-1…4、METEOR、ROUGE-L 与 CIDEr-D。")

    heading(doc, "2.4 结果怎样理解", 2)
    for text, lead in (
        ("完整 MCINet 报告 mIoU 86.43、BLEU-4 65.95、CIDEr-D 140.29。数值的关键意义不是绝对领先幅度，而是同一视觉底座同时支撑像素与语言输出。", "总体结果："),
        ("多任务相对单任务：检测单任务 mIoU 86.54，联合后 86.43，检测略降；描述则从 BLEU-4 65.86 / CIDEr 139.92 小幅升到 65.95 / 140.29。说明任务协同存在，但不是所有指标都单调受益。", "多任务关系："),
        ("LPE/GDFA 对描述增益更明显，mIoU 改善有限；小于 400 像素的建筑 IoU 即使在完整模型中也只有 17.24，暴露了低分辨率小目标的主要瓶颈。", "模块作用："),
        ("论文自己承认任务平衡仍困难：描述可能忽略次要变化，反过来削弱检测；Agent 还依赖手工 prompt、工具调度与有限工具集。", "限制："),
    ):
        list_item(doc, bullet_id, text, lead)

    lead_callout(doc, "方法评价", "MCINet 的研究贡献是“真正共享、多尺度、双输出”的 MCI 体系；Agent 是其交互层扩展。对于当前项目，最值得借鉴的是多尺度视觉交互、损失动态归一化和小目标分层评估，而不是简单把一个低分辨率 mask head 视作等价实现。")

    heading(doc, "3. SECOND-CC 论文分析：数据鲁棒性与 MModalCC", 1)
    body(doc, "论文：Ali Can Karaca 等，Robust Change Captioning in Remote Sensing: SECOND-CC Dataset and MModalCC Framework，IEEE JSTARS 2025。论文同时关注真实遥感图中的错位、模糊、亮度和分辨率变化，并利用前后语义图提高描述鲁棒性。")
    add_link_line(doc, [
        ("论文原文", "https://arxiv.org/html/2501.10075v1"),
        ("DOI", "https://doi.org/10.1109/JSTARS.2025.3600613"),
        ("官方仓库", "https://github.com/ChangeCapsInRS/SecondCC"),
    ])

    heading(doc, "3.1 数据集设计", 2)
    body(doc, "SECOND-CC 来源于 SECOND 语义变化检测数据，覆盖杭州、成都、上海。原始 512×512 影像被切成四个 256×256 子块，空间分辨率约 0.5–3 m/pixel。每个图像对有前后 RGB、前后土地覆盖语义图和人工描述；七名标注者历时约一年，强调变化类型、位置、强度与多变化顺序，并规范方向词。")
    body(doc, "论文将真实干扰视为数据本身的一部分，并在 train/val 上增加高斯模糊、亮度变化、镜像与旋转。模糊/亮度增强可复制描述；几何增强必须同步改写 left/right/top/bottom 等方向词。这是本地 AUG 版本的核心价值，也是需要检查的潜在标签噪声来源。")

    heading(doc, "3.2 MModalCC 方法论拆解", 2)
    flow_box(doc, "RGB(T0,T1) + Semantic(T0,T1) → 双 Siamese ResNet101 → CMCA ↔ UDCA ×2 → MGCA 门控解码 → caption")
    for text, lead in (
        ("双模态编码：RGB 对和语义图对分别由独立的 Siamese ResNet101 编码；每一模态内部两时相共享权重，但 RGB 与语义不共享编码器。", "双模态编码："),
        ("CMCA：在同一时相内做跨模态注意，使 RGB 纹理与语义类别互相校正。", "CMCA："),
        ("UDCA：在每个模态内部利用时相差异作为引导，聚焦真正变化并抑制成像扰动。CMCA 与 UDCA 交替两轮，且共享跨注意参数，形成迭代特征增强。", "UDCA："),
        ("MGCA 解码：文本 query 分别对 RGB 与 semantic 特征 cross-attend，再用 sigmoid 门控自适应融合文本状态、RGB 上下文和语义上下文，最后由 Transformer 自回归生成描述。", "MGCA 解码："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "3.3 训练、指标与结果", 2)
    body(doc, "论文采用 ImageNet 预训练 ResNet101 并 fine-tune，Adam 学习率 5e-5，30 epochs，交叉熵训练。论文表中选择 beam size 2；官方 README 的示例推理写 beam 4，因此复现时应以论文实验表和发布 checkpoint 的实际脚本为准。评测包含 BLEU-1…4、METEOR、ROUGE-L、CIDEr-D、SPICE，并分别报告 change/no-change。no-change 的参考句高度相同，CIDEr 会退化为 0，论文不把它作为 no-change 的核心指标。")
    for text, lead in (
        ("MModalCC 总体报告 BLEU-4 0.386、CIDEr-D 0.933、SPICE 0.249。RGB+semantic 明显优于仅 RGB 或仅 semantic，说明互补信息确实进入了语言解码。", "总体结果："),
        ("AUG 将总体 BLEU-4 从 0.383 提至 0.386、CIDEr 从 0.898 提至 0.933；change 子集 BLEU-4 从 0.240 提至 0.261。增益主要来自更困难的变化样本，而非简单记忆 no-change 模板。", "增强作用："),
        ("CMCA+UDCA 优于任一单独模块；双模态解码也优于单模态。但论文对部分基线只提供单模态输入，而 MModalCC 同时使用两种模态，比较时必须把输入信息差异写清楚。", "模块作用与公平性："),
        ("语义图来自真值标注。真实部署中需要预测语义图，预测错误如何传导到 caption 没有被系统评估；这是论文和当前代码共同的 oracle-semantic 限制。", "限制："),
    ):
        list_item(doc, bullet_id, text, lead)

    lead_callout(doc, "方法评价", "MModalCC 的关键不是“加一层语义 attention”，而是双编码器、跨模态/跨时间迭代增强与解码端门控三者协同。当前代码的轻量融合可以作为有效消融，但不能命名为 MModalCC 复现。")

    heading(doc, "4. 当前实验代码的方法画像", 1)
    heading(doc, "4.1 CARD 主干", 2)
    flow_box(doc, "离线 ResNet101 stage-3 特征（1024×14×14） → CARD 公共/差异分解 → DynamicSpeaker Transformer → caption")
    body(doc, "输入不是原始 RGB，而是先把图像 bicubic 缩放到 224×224，再用 ImageNet ResNet101 stage-3 离线提取约 1024×14×14 的 .npy 特征。训练时视觉骨干被冻结且没有在线图像增强。CARD 对两个时相分别做一次自注意，用 in-batch 对比损失约束公共表示，用 HSIC 约束差异表示的独立性；随后做跨时相注意、从原特征中减去公共上下文，拼接双向差异 tokens，交给 2 层、8 头、512 维 DynamicSpeaker 解码。")
    body(doc, "视觉配置中的 transformer_encoder.att_layer=2 仅被保存为成员变量，当前 CARD 视觉侧没有按该值循环堆叠；真正的堆叠层数只明确出现在语言解码器。这个细节会影响对“2 层视觉 Transformer”的描述。")

    heading(doc, "4.2 LEVIR-MCI 最终配置", 2)
    for text, lead in (
        ("三类 mask 辅助：从 CARD 差异 tokens 恢复 14×14 空间特征，用 CE+Dice 预测背景/道路/建筑；λmask=0.003。", "三类 mask 辅助："),
        ("文本语义辅助：从每图全部参考 caption 按规则抽取 21 维 object/action multi-label 标签，用 BCE 训练全局语义头；λsem=0.005。它不是像素级语义标注。", "文本语义辅助："),
        ("弱耦合方式：semantic partial detach=0.5；mask reweight、semantic cross-attention 和 hard gate 均关闭。因此 mask/tag 只在训练期塑造视觉表示，不在推理时直接进入 caption 输入。", "弱耦合方式："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "4.3 SECOND-CC 最终配置", 2)
    for text, lead in (
        ("真值语义输入：训练、验证和测试都读取 sem/A 与 sem/B。模型将前后类别嵌入与绝对差异拼接为 key/value，以 CARD 差异 tokens 为 query，做一次 MultiheadAttention，并用可学习 gamma 残差融合。", "真值语义输入："),
        ("语义差异辅助：未变化像素为 0，变化像素为 after_class+1；14×14 dense head 以 CE+Dice 训练，λsem=0.005。", "语义差异辅助："),
        ("方法性质：这是“RGB CARD + oracle semantic maps”的多模态上界设置，不是只在训练期用语义监督的 RGB-only 模型。正式 Test 也把真值前后语义图送入 caption 分支。", "方法性质："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "4.4 训练目标与选择协议", 2)
    body(doc, "实际总损失为 Lcaption + 0.001 Lcommon + 0.001 LHSIC + λmask Lmask + λsem Lsemantic。两个最终配置均使用 batch 32、10,000 iterations、Adam 2e-4、StepLR(step_size=17, gamma=0.1)、seed 1111，每 1,000 step 保存/验证，未启用梯度裁剪。辅助权重在前 30% 为 0，30%–70% 线性升温；semantic 在后段衰减至目标值的 50%。")
    body(doc, "仓库提供 validation-only selector、manifest 锁定和 change/no-change 汇总工具，这是良好的论文工程基础；但 train_card_spot.py 内部的 selection_strategy 主要用于记录，训练期 best_balanced 仍依赖一组未按数据集区分的硬编码基线。正式结果应以外部 selector 产生的 validation manifest 为唯一选择依据。")

    heading(doc, "5. 论文方法与当前代码的逐项对比", 1)
    heading(doc, "5.1 LEVIR-MCI：MCINet vs 当前 CARD 扩展", 2)
    add_table(
        doc,
        ["维度", "论文 MCINet / Change-Agent", "当前 LEVIR-MCI 代码", "评价与影响"],
        [
            ["视觉输入", "原图、多尺度、端到端 SegFormer-B1", "冻结的 ResNet101 stage-3 14×14 特征", "计算更轻、控制变量清晰；但损失空间细节与小目标能力"],
            ["时相交互", "BI3：LPE + GDFA，多层/多尺度", "CARD 公共/差异分解 + 一次跨时相注意", "研究假设不同，不能视为同一模块替代"],
            ["检测分支", "CBF 多尺度融合并恢复像素级掩膜", "两层卷积式辅助头，输出 14×14", "适合辅助监督；不等价于论文检测能力"],
            ["语言分支", "BI3 后接 Transformer decoder", "CARD tokens + DynamicSpeaker", "保留 CARD 的强项，但未利用论文的差异引导"],
            ["任务耦合", "共享骨干 + 自归一化双任务损失", "固定 λ + warmup + partial detach", "当前耦合更保守；应补 loss normalization 消融"],
            ["推理输出", "mask + caption；Agent 可进一步编排", "caption；辅助 head 可监控，未实现 Agent", "应明确研究范围，不使用 Change-Agent 名称"],
            ["指标", "原分辨率 mIoU + caption；小目标分析", "14×14 宏平均辅助 IoU + caption", "分割数字不可直接横比"],
        ],
        [1120, 2570, 2580, 3090],
        font_size=8.2,
    )
    lead_callout(doc, "定位结论", "当前实现最合理的论文定位是：在固定 CARD 表征上研究像素掩膜与文本语义辅助是否能改善变化描述。它不是 MCINet 复现；若论文声称“多层次变化解释”，至少需要加入官方 MCINet 外部基线、原分辨率 mask 指标和小目标分析。")

    heading(doc, "5.2 SECOND-CC：MModalCC vs 当前 CARD 扩展", 2)
    add_table(
        doc,
        ["维度", "论文 MModalCC", "当前 SECOND-CC 代码", "评价与影响"],
        [
            ["视觉编码", "RGB 与 semantic 各一套 Siamese ResNet101，可 fine-tune", "RGB 为冻结离线特征；semantic 为类别 embedding", "轻量但模态表征能力不对等"],
            ["模态交互", "CMCA 与 UDCA 交替两轮，跨模态+跨时间", "一次 semantic K/V → CARD query cross-attention", "可做轻量替代消融，不能等同 MModalCC"],
            ["解码融合", "MGCA 分别读取 RGB/semantic，并门控三路信息", "融合先发生在 caption tokens，普通 DynamicSpeaker 解码", "缺少词级动态模态选择"],
            ["语义监督", "真值前后语义图作为输入", "同样在 Test 使用真值前后语义图；另预测 14×14 diff", "比较时必须标注 oracle semantic input"],
            ["增强", "train/val 静态增强，几何方向词同步修改", "直接使用官方/本地 AUG 文件", "协议基本一致；需检查描述缺失与方向词质量"],
            ["生成", "论文 beam=2（README 示例 beam=4）", "验证随机 sampling；Test greedy；beam_size 未使用", "当前协议不一致，是 P0 修复项"],
            ["评测", "change/no-change 分组；no-change CIDEr 非核心", "具备分组工具，但需单独运行", "应纳入自动锁定 Test 流程"],
        ],
        [1120, 2570, 2580, 3090],
        font_size=8.2,
    )
    lead_callout(doc, "定位结论", "当前语义 cross-attention 是一条有价值的低成本路线：它直接检验“oracle semantic maps 对 CARD captioning 的增益”。但只要 Test 仍读取 GT sem/A,B，结果就只能与同样使用真值语义图的方法比较，不能作为 RGB-only 部署能力。")

    heading(doc, "6. 代码审计发现：影响论文结论的风险", 1)
    add_table(
        doc,
        ["优先级", "发现", "代码证据", "影响", "建议"],
        [
            [("P0", True), "验证默认 multinomial sampling；Test 显式 greedy", "transformer_decoder.py:227–257；train_card_spot.py:1566；test_card_spot.py:461", "checkpoint 排名随机波动，选择协议与最终测试不一致", "验证也显式 sample_max=1，并统一单一 decode 函数"],
            [("P0", True), "SECOND 基础语义映射只产生 0…5，配置写 7，Dataset 再 +1", "rcc_dataset_transformer_levir.py:39–53, 165–170, 453–458；second YAML:11,51", "最终 8-logit dense head 至少有一类永无正样本；embedding 也过配", "核对 taxonomy；大概率 base=6、diff=7，并加类别覆盖单测"],
            [("P0", True), "本地没有 features/ 与任何 .pt/.pth checkpoint", "Dataset preflight:262–315；本地文件盘点", "当前无法训练/复核历史结果；正式清单为 missing/failed", "先生成全部 ResNet 特征、校验数量/shape/hash，再启动正式实验"],
            [("P0", True), "SECOND Test 读取 GT sem/A,B", "test_card_spot.py:439–445；CARD.py:554–567", "输入信息高于 RGB-only 基线", "结果命名为 RGB+OracleSem；补 RGB-only 与 predicted/noisy-sem"],
            [("P1", True), "分割在 14×14 上评测，GT 下采样；再做每图平均", "test_card_spot.py:67–135", "不能对照论文原分辨率全局 mIoU；no-change 正确全背景仍记 0", "用全局 confusion matrix；另报原分辨率上采样/高分辨率 head"],
            [("P1", True), "训练内 selection_strategy 不真正切换选择算法；balanced 常量不分数据集", "train_card_spot.py:41–47；外部 select_best_checkpoint.py", "配置名与实际选择可能错位", "只允许 validation selector manifest 进入 Test，并审计数据集基线"],
            [("P1", True), "配置 beam_size 未被 DynamicSpeaker 使用", "transformer_decoder.py:227–257；test_card_spot.py:461", "与论文 beam 协议不一致；搜索收益未知", "实现/验证 beam，或明确所有方法统一 greedy"],
            [("P2", True), "视觉 att_layer=2 未实际堆叠；位置参数固定 14²+1", "CARD.py:346,374,467–468,513–514", "配置表意偏差；更换分辨率/backbone 易失败", "删除误导参数或真正循环堆叠，并动态位置编码"],
        ],
        [780, 2130, 2250, 2000, 2200],
        font_size=7.8,
    )

    lead_callout(doc, "当前可复现性评级", "方法审计：可；配置审计：可；完整训练：不可（缺 features）；历史性能复核：不可（缺 checkpoint）；论文级结论：暂不可。评级针对当前工作区状态，不否定远程机器上可能存在的工件。", fill=PALE_RED, left_color="B42318")

    heading(doc, "6.1 还需注意的次级问题", 2)
    for text, lead in (
        ("HSIC 分母为 (batch_size−1)²，batch=1 的 smoke test 会产生无效值。", "数值稳定性："),
        ("feature preflight 只检查每个 split 的首样本，缺文件可能在训练中途才暴露。", "输入完整性："),
        ("allow_missing_pseudo_mask=False 仍可能只是警告并填充 ignore target，不会严格失败。", "掩膜完整性："),
        ("requirements.txt 的 torch wheel 与 COCO scorer 路径带有原作者机器痕迹，环境复现仍需清理。", "环境可移植性："),
        ("checkpoint 只保存 CARD、speaker 和 model_cfg，不支持优化器/调度器完整 resume。", "长训练恢复："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "7. 意见与评价", 1)
    heading(doc, "7.1 值得保留的研究设计", 2)
    for text, lead in (
        ("统一主干的控制实验：两个新数据集都落在同一 CARD 主干上，便于回答“额外监督是否有效”，比同时改 backbone、decoder、loss 更容易归因。", "统一主干的控制实验："),
        ("弱耦合与梯度控制：warmup、partial detach、可学习 gamma 都是在面对 noisy auxiliary signal 时合理的稳定手段。", "弱耦合与梯度控制："),
        ("论文工程意识：仓库已经具备 validation-only selector、manifest、锁定 Test、结构化指标与 change/no-change 分组工具，方向正确。", "论文工程意识："),
        ("数据协议意识：SECOND 的增强只作用于 train/val、test 保持原始，符合论文设置；LEVIR 的变化/无变化也近似均衡。", "数据协议意识："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "7.2 当前方法的主要短板", 2)
    for text, lead in (
        ("空间信息瓶颈：离线 14×14 特征和低分辨率辅助 head 天然不利于 LEVIR-MCI 的小道路、小建筑与精确边界。", "空间信息瓶颈："),
        ("方法命名边界：LEVIR 路线不具备 BI3/CBF/Agent，SECOND 路线不具备双语义编码器、CMCA/UDCA/MGCA；只能称 CARD-based extension。", "方法命名边界："),
        ("输入公平性：SECOND 目前是 oracle-semantic 上界。若不同时给 baseline 相同语义输入，绝对分数比较会误导。", "输入公平性："),
        ("实验状态不足：没有可运行特征、没有当前 checkpoint、正式表为空，多随机种子结论尚未成立。", "实验状态不足："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "7.3 总体评价", 2)
    add_table(
        doc,
        ["维度", "当前评级", "依据"],
        [
            ["研究问题与原创性", "4 / 5", "统一 CARD 主干比较多粒度监督，问题清晰；若补 predicted prior 与鲁棒性曲线，辨识度更强。"],
            ["与前沿方法契合度", "4 / 5", "公共/差异分解与 2026 的 prior/语义对齐方向兼容，但尚缺层次路由和多尺度在线视觉。"],
            ["实验协议可信度", "2 / 5", "验证随机解码、Test greedy、SECOND 类别口径与 oracle-sem 输入尚未修正/分栏。"],
            ["可复现性", "2 / 5", "代码和配置可审计，但本地无 features/checkpoint，正式实验记录仍 missing/failed。"],
            ["部署可行性", "2 / 5", "冻结离线特征较省训练，但依赖离线预处理；SECOND 依赖 GT 语义图，且尚无吞吐/延迟报告。"],
            ["当前论文就绪度", "2 / 5", "可以写方法与研究问题，尚不能写可靠性能结论；完成 P0、三随机种子和公平基线后再评。"],
        ],
        [2350, 1350, 5660],
        font_size=8.3,
    )
    lead_callout(
        doc,
        "总体判断",
        "当前项目具有清晰而有价值的研究主线：用 mask/semantic supervision 弱耦合地增强 CARD，而不是重做两个论文模型。代码已经从“想法原型”进入“可审计的实验基础设施”阶段，但还没有跨过“可复现、可公平比较、可统计下结论”的门槛。最重要的不是继续搜索 λ，而是先修正协议与类别定义，再补齐工件和官方基线。",
    )

    heading(doc, "8. 建议的修复顺序与最小实验矩阵", 1)
    heading(doc, "8.1 P0：先让协议可信、实验可运行", 2)
    for text in (
        "把验证和测试统一为 deterministic greedy（或统一 beam），固定 seed，并把 decode 配置写入每个结果 JSON。",
        "修正 SECOND 类别定义：先对所有 sem/A,B 扫描颜色与 class id 覆盖；若基础类确为 0…5，则配置 base=6、diff=7。添加断言：每个 logit 类在 train 至少出现一次。",
        "运行官方特征提取脚本生成两数据集 A/B 的 ResNet101 stage-3 特征；全量校验文件数、shape、dtype、NaN 和 annotation 对齐，而不是只检查首样本。",
        "只允许外部 validation selector 生成的 manifest 进入 Test；manifest 保存 config hash、checkpoint hash、decode、dataset version 与语义输入模式。",
        "重写分割评测：累积全局 confusion matrix；no-change 全背景正确时单独记 specificity/accuracy，不把 0/0 强制当 IoU=0。",
    ):
        list_item(doc, number_id, text)

    heading(doc, "8.2 P1：建立公平的主结果表", 2)
    add_table(
        doc,
        ["数据集", "最小对照组", "必须控制的变量", "核心报告"],
        [
            ["LEVIR-MCI", "CARD RGB；+mask；+caption-tags；+mask+tags；+paper loss normalization；官方 MCINet；LCCC", "相同 split、相同解码、相同 seed 集；区分冻结特征与端到端；LCCC 需联合 mask+caption", "8 项 caption；原分辨率/全局 mask mIoU；road/building IoU；小目标 IoU；均值±标准差"],
            ["SECOND-CC-AUG", "CARD RGB；+dense aux（训练期）；+oracle-sem cross-attn；RGB+noisy/predicted-sem；官方 MModalCC", "统一 AUG、同输入模态、同解码；明确 oracle vs predicted", "总体 + change/no-change 8 项 caption；no-change CIDEr 仅附录；参数量/FLOPs/延迟"],
        ],
        [1370, 3010, 2350, 2630],
        font_size=8.2,
    )

    heading(doc, "8.3 P2：最有信息增益的扩展实验", 2)
    for text, lead in (
        ("LEVIR 多尺度：在保持 CARD decoder 不变的前提下，引入 stage2+stage3 特征或轻量上采样 mask decoder；按目标面积分桶，验证提升是否集中在小建筑。", "LEVIR 多尺度："),
        ("任务平衡：把当前 fixed λ+warmup 与 MCINet 的 stop-gradient loss normalization 并列，记录每个任务的梯度范数与 caption/mIoU Pareto 前沿。", "任务平衡："),
        ("SECOND 语义噪声：随机翻转 5%/10% 类别、边界腐蚀/膨胀、用预测语义图替代 GT，画出 caption 指标随语义质量下降的曲线。", "SECOND 语义噪声："),
        ("解码融合：在现有单层 cross-attention 之外，做词级门控 ablation，验证是否真的需要 MModalCC 风格 MGCA，而不是无条件添加语义上下文。", "解码融合："),
        ("效率：当前轻量方法的潜在优势是低计算成本。与官方模型对比时同时报告参数量、训练显存、吞吐和单图延迟，才能把“轻量替代”写成贡献。", "效率："),
    ):
        list_item(doc, bullet_id, text, lead)

    heading(doc, "9. 论文写作时建议采用的结论边界", 1)
    heading(doc, "9.1 可以成立的表述", 2)
    for text in (
        "“我们在统一 CARD 主干上研究像素级变化掩膜、caption 派生语义标签与前后语义图对变化描述的影响。”",
        "“LEVIR-MCI 路线使用训练期辅助监督；SECOND 路线使用 RGB + oracle semantic maps，二者信息条件不同。”",
        "“当前方法是轻量 CARD-based extension；MCINet 与 MModalCC 作为外部专用架构基线。”",
        "“所有 checkpoint 仅依据 validation 选择，Test 在 manifest 锁定后运行一次；报告多随机种子均值和标准差。”",
    ):
        list_item(doc, bullet_id, text)

    heading(doc, "9.2 在修复/补实验前不应使用的表述", 2)
    for text in (
        "“复现了 Change-Agent/MCINet”或“复现了 MModalCC”。",
        "把 14×14 辅助 mIoU 与 MCINet 原分辨率 mIoU 放在同一列直接比较。",
        "把 SECOND 的 oracle-semantic 分数与 RGB-only baseline 直接解释为架构更优。",
        "依据当前工作区的历史 Markdown 单次 Test 数值宣称稳定提升。",
        "把 selection_strategy 配置名当作实际执行过的 checkpoint 选择证据。",
    ):
        list_item(doc, bullet_id, text)

    lead_callout(doc, "推荐论文主张", "最稳妥也最有辨识度的贡献主张是：在 CARD 公共/差异表征上，以可控梯度的弱耦合方式注入不同粒度的变化监督，并系统研究监督粒度、输入模态与鲁棒性的权衡。这个主张比“复现两个专用模型”更贴合现有代码。")

    heading(doc, "10. 本地代码证据索引", 1)
    add_table(
        doc,
        ["主题", "文件与行号", "证据"],
        [
            ["CARD 主干", "models/CARD.py:444–527", "公共/差异分解、对比约束、HSIC、跨时相注意与差异 tokens"],
            ["语义融合", "models/CARD.py:203–298, 544–594", "前后语义 embedding、abs diff、K/V cross-attention、gamma 残差、partial detach"],
            ["总损失", "train_card_spot.py:1204–1348", "warmup 与 caption/common/HSIC/mask/semantic 加权"],
            ["验证/测试解码", "train_card_spot.py:1566；models/transformer_decoder.py:227–257；test_card_spot.py:461", "validation sampling vs Test greedy"],
            ["SECOND 类别", "datasets/rcc_dataset_transformer_levir.py:39–53,165–170,453–458", "0…5 映射、基础类+1、after+1 的 diff target"],
            ["低分辨率指标", "test_card_spot.py:67–135", "target 下采样到 logits、每图/批次平均、空 union 处理"],
            ["离线特征", "scripts/extract_features.py:40–75,113–135；datasets/...:262–315,552", "ImageNet ResNet101 stage-3 与 .npy preflight/读取"],
            ["正式实验状态", "experiments/paper_required_experiments_summary.csv:2–10", "相关训练 missing/failed；外部 MModalCC pending"],
        ],
        [1700, 3540, 4120],
        font_size=8.3,
    )

    heading(doc, "参考文献与链接", 1)

    heading(doc, "2025—2026 最新与前沿工作", 2)
    add_reference(doc, "R1", "Zhang, L.; Lam, S.-K.; Akhtar, N. LBTCap: A Lightweight Bilateral Transformer for Real-Time Remote Sensing Image Change Captioning. arXiv, 2026-07-03.", [("arXiv", "https://arxiv.org/abs/2607.03320"), ("HTML", "https://arxiv.org/html/2607.03320")])
    add_reference(doc, "R2", "Liu, Z. et al. JL1-CC&QA: Extending the JL1-CD Benchmark with Change Captioning and Question Answering. arXiv, 2026-06-30.", [("arXiv", "https://arxiv.org/abs/2606.31745"), ("数据仓库", "https://github.com/circleLZY/JL1-CD")])
    add_reference(doc, "R3", "Wang, Y. et al. RSICCLLM: A Multimodal Large Language Model for Remote Sensing Image Change Captioning. ECCV 2026 accepted; arXiv:2606.28266.", [("arXiv", "https://arxiv.org/abs/2606.28266"), ("项目仓库", "https://github.com/keaill/RSICCLLM")])
    add_reference(doc, "R4", "Wang, Y. et al. DFM: Difference Feature Modeling with Text-Guided Gated Contrastive Loss for Remote Sensing Image Change Captioning. IEEE ICME 2026 accepted; arXiv:2606.27410.", [("arXiv", "https://arxiv.org/abs/2606.27410")])
    add_reference(doc, "R5", "HiSem: Hierarchical Semantic Disentangling for Remote Sensing Image Change Captioning. IEEE TGRS, 64, 2026.", [("DOI", "https://doi.org/10.1109/TGRS.2026.3705251"), ("arXiv", "https://arxiv.org/abs/2605.15024"), ("代码入口", "https://github.com/Man-Wang-star/HiSem")])
    add_reference(doc, "R6", "UAV as Urban Construction Change Monitor: A New Benchmark and Change Captioning Model. arXiv:2605.04409, 2026.", [("arXiv", "https://arxiv.org/abs/2605.04409"), ("代码（数据待开放）", "https://github.com/G124556/ptnet")])
    add_reference(doc, "R7", "STAND: Semantic Anchoring Constraint with Dual-Granularity Disambiguation for Remote Sensing Image Change Captioning. arXiv:2604.23309, 2026.", [("arXiv", "https://arxiv.org/abs/2604.23309")])
    add_reference(doc, "R8", "Zhao, Y. et al. Disturbance-Robust Remote Sensing Change Captioning with Self-supervised Multifrequency Representation. Journal of Remote Sensing, 6:1037, 2026.", [("DOI", "https://doi.org/10.34133/remotesensing.1037")])
    add_reference(doc, "R9", "Li, Y. et al. Exploring Difference Semantic Prior Guidance for Remote Sensing Image Change Captioning. Remote Sensing, 18(2):232, 2026.", [("DOI", "https://doi.org/10.3390/rs18020232")])
    add_reference(doc, "R10", "Describing Land Cover Changes via Multi-Temporal Remote Sensing Image Captioning Using LLM, ViT, and LoRA. Remote Sensing, 18(1):166, 2026.", [("DOI", "https://doi.org/10.3390/rs18010166")])
    add_reference(doc, "R11", "Ricci, R.; Bazi, Y.; Melgani, F. Remote sensing change captioning meets large language and vision models. ISPRS Journal of Photogrammetry and Remote Sensing, online first, 2026.", [("DOI", "https://doi.org/10.1016/j.isprsjprs.2026.06.003")])
    add_reference(doc, "R12", "Yang, C. et al. Enhancing Perception of Key Changes in Remote Sensing Image Change Captioning. IEEE Transactions on Image Processing, 34:7378–7390, 2025.", [("DOI", "https://doi.org/10.1109/TIP.2025.3589096"), ("GitHub", "https://github.com/yangcong356/KCFI")])
    add_reference(doc, "R13", "Change Captioning in Remote Sensing: Evolution to SAT-Cap — A Single-Stage Transformer Approach. arXiv:2501.08114, 2025.", [("arXiv", "https://arxiv.org/abs/2501.08114")])
    add_reference(doc, "R14", "Chen, Z. et al. RSCC: A Large-Scale Remote Sensing Change Caption Dataset for Disaster Events. NeurIPS 2025 Datasets & Benchmarks Track.", [("arXiv", "https://arxiv.org/abs/2509.01907"), ("项目页", "https://bili-sakura.github.io/RSCC/"), ("GitHub", "https://github.com/Bili-Sakura/RSCC")])
    add_reference(doc, "R15", "Towards Comprehensive Interactive Change Understanding in Remote Sensing: A Large-scale Dataset and Dual-granularity Enhanced VLM. arXiv:2509.23105, 2025.", [("arXiv", "https://arxiv.org/abs/2509.23105")])
    add_reference(doc, "R16", "SAM Guided Semantic and Motion Changed Region Mining for Remote Sensing Change Captioning. IEEE TGRS, 64, 2026.", [("DOI", "https://doi.org/10.1109/TGRS.2026.3709816"), ("arXiv", "https://arxiv.org/abs/2511.21420")])
    add_reference(doc, "R17", "Li, H.; Zhang, Y.; Tu, Y.; Zhang, Y.; Wang, W.; Li, L. Learning to Change by Critique and Correction: A Synergistic Framework for Remote Sensing Change Detection and Captioning. IEEE TIP, 35:7783–7798, 2026.", [("DOI", "https://doi.org/10.1109/TIP.2026.3713472"), ("GitHub", "https://github.com/Throb16/Lccc")])
    add_reference(doc, "R18", "Zhou, Q. et al. Weakly Paired Remote Sensing Change Captioning With Large Multimodal Model. IEEE TGRS, 64, 2026.", [("DOI", "https://doi.org/10.1109/TGRS.2026.3698503")])
    add_reference(doc, "R19", "Deng, P.; Zhou, W.; Wu, H. DeltaVLM: Interactive Remote Sensing Image Change Analysis via Instruction-Guided Difference Perception. Remote Sensing, 18(4):541, 2026.", [("DOI", "https://doi.org/10.3390/rs18040541"), ("arXiv", "https://arxiv.org/abs/2507.22346"), ("GitHub", "https://github.com/hanlinwu/DeltaVLM")])

    heading(doc, "基础方法脉络", 2)
    add_reference(doc, "F1", "Chouaf, S. et al. Captioning Changes in Bi-Temporal Remote Sensing Images. IGARSS 2021.", [("DOI", "https://doi.org/10.1109/IGARSS47720.2021.9554419")])
    add_reference(doc, "F2", "Hoxha, G. et al. Change Captioning: A New Paradigm for Multitemporal Remote Sensing Image Analysis. IEEE TGRS, 60, 2022.", [("DOI", "https://doi.org/10.1109/TGRS.2022.3195692")])
    add_reference(doc, "F3", "Liu, C. et al. Remote Sensing Image Change Captioning With Dual-Branch Transformers: A New Method and a Large Scale Dataset. IEEE TGRS, 60, 2022.", [("DOI", "https://doi.org/10.1109/TGRS.2022.3218921"), ("GitHub", "https://github.com/Chen-Yang-Liu/RSICC")])
    add_reference(doc, "F4", "Liu, C. et al. Progressive Scale-Aware Network for Remote Sensing Image Change Captioning. IGARSS 2023.", [("DOI", "https://doi.org/10.1109/IGARSS52108.2023.10283451"), ("GitHub", "https://github.com/Chen-Yang-Liu/PSNet")])
    add_reference(doc, "F5", "Liu, C. et al. A Decoupling Paradigm With Prompt Learning for Remote Sensing Image Change Captioning. IEEE TGRS, 61, 2023.", [("DOI", "https://doi.org/10.1109/TGRS.2023.3321752"), ("GitHub", "https://github.com/Chen-Yang-Liu/PromptCC")])
    add_reference(doc, "F6", "Chang, S.; Ghamisi, P. Changes to Captions: An Attentive Network for Remote Sensing Change Captioning. IEEE TIP, 32, 2023.", [("DOI", "https://doi.org/10.1109/TIP.2023.3328224"), ("GitHub", "https://github.com/ShizhenChang/Chg2Cap")])
    add_reference(doc, "F7", "Liu, C. et al. RSCaMa: Remote Sensing Image Change Captioning With State Space Model. IEEE GRSL, 21, 2024.", [("DOI", "https://doi.org/10.1109/LGRS.2024.3404604"), ("GitHub", "https://github.com/Chen-Yang-Liu/RSCaMa")])
    add_reference(doc, "F8", "Zhu, R. et al. Semantic-CC: Boosting Remote Sensing Image Change Captioning via Foundational Knowledge and Semantic Guidance. IEEE TGRS, 62, 2024.", [("DOI", "https://doi.org/10.1109/TGRS.2024.3497338"), ("arXiv", "https://arxiv.org/abs/2407.14032")])

    heading(doc, "两个新数据集与当前主干", 2)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    set_run_font(p.add_run("[1] "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run("Liu, C. et al. Change-Agent: Toward Interactive Comprehensive Remote Sensing Change Interpretation and Analysis. IEEE Transactions on Geoscience and Remote Sensing, 2024. "))
    add_hyperlink(p, "arXiv", "https://arxiv.org/abs/2403.19646")
    set_run_font(p.add_run(" · "), color=MUTED)
    add_hyperlink(p, "DOI", "https://doi.org/10.1109/TGRS.2024.3425815")
    set_run_font(p.add_run(" · "), color=MUTED)
    add_hyperlink(p, "GitHub", "https://github.com/Chen-Yang-Liu/Change-Agent")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    set_run_font(p.add_run("[2] "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run("Karaca, A. C. et al. Robust Change Captioning in Remote Sensing: SECOND-CC Dataset and MModalCC Framework. IEEE Journal of Selected Topics in Applied Earth Observations and Remote Sensing, 2025. "))
    add_hyperlink(p, "arXiv", "https://arxiv.org/abs/2501.10075")
    set_run_font(p.add_run(" · "), color=MUTED)
    add_hyperlink(p, "DOI", "https://doi.org/10.1109/JSTARS.2025.3600613")
    set_run_font(p.add_run(" · "), color=MUTED)
    add_hyperlink(p, "GitHub", "https://github.com/ChangeCapsInRS/SecondCC")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    set_run_font(p.add_run("[3] "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run("Tu, Y. et al. Context-aware Difference Distilling for Multi-change Captioning. ACL 2024, pp. 7941–7956. "))
    add_hyperlink(p, "ACL Anthology", "https://aclanthology.org/2024.acl-long.430/")
    set_run_font(p.add_run(" · "), color=MUTED)
    add_hyperlink(p, "PDF", "https://aclanthology.org/2024.acl-long.430.pdf")

    heading(doc, "分析说明", 2)
    body(doc, "论文方法、训练与公开结果均以论文原文/官方仓库为准；本地实现判断来自 2026-07-26 的 c906eb3 快照。对论文模块作用的表述属于基于论文消融结果的解释；对当前代码风险的表述来自源码静态审计与本地工件盘点，未执行 GPU 训练。")
    body(doc, "文档样式采用 standard_business_brief 预设；中文字体仅做 Microsoft YaHei 东亚字形替代，西文字体保持 Calibri。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
